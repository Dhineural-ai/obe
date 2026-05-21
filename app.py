import streamlit as st
import PyPDF2
import re
import os

from io import BytesIO
from dotenv import load_dotenv
from openai import OpenAI
from docx import Document

load_dotenv()

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY")
)

st.set_page_config(
    page_title="AI Curriculum Intelligence Platform",
    layout="wide"
)

st.title("AI Powered Curriculum Intelligence Platform")

st.markdown("""
This platform performs:
- Curriculum Intelligence Audit
- OBE Redesign
- CO Redesign
- Intelligent CO-PO Mapping
- Pedagogy Intelligence
- Assignment Intelligence
- Industry Alignment
- Accreditation Readiness Enhancement
""")

vision_file = st.file_uploader(
    "Upload Vision & Mission PDF",
    type=["pdf"]
)

peo_file = st.file_uploader(
    "Upload PEO PO PSO PDF",
    type=["pdf"]
)

course_file = st.file_uploader(
    "Upload Course File PDF",
    type=["pdf"]
)

def extract_text_from_pdf(pdf_file):

    text = ""

    pdf_reader = PyPDF2.PdfReader(pdf_file)

    for page in pdf_reader.pages:

        extracted = page.extract_text()

        if extracted:

            extracted = extracted.replace("\n", " ")

            extracted = re.sub(r"\s+", " ", extracted)

            text += extracted + "\n\n"

    return text

def generate_curriculum_intelligence_report(
    vision_text,
    peo_text,
    course_text
):

    prompt = f"""

You are:
- NBA expert
- NAAC expert
- NIRF consultant
- Curriculum architect
- Bloom taxonomy specialist
- OBE strategist

IMPORTANT:
DO NOT trust uploaded syllabus blindly.

Assume:
- prerequisites may be weak
- COs may be weak
- mappings may be illogical
- pedagogy may be ineffective

unless proven otherwise.

DO NOT simply paraphrase.

You must critically audit and redesign curriculum according to:
- global standards
- NBA expectations
- graduate competency expectations
- industry relevance
- Bloom taxonomy
- employability
- practical competency

DO NOT CHANGE:
- unit titles
- unit contents

You ARE allowed to redesign:
- prerequisites
- objectives
- course outcomes
- mappings
- pedagogy
- assignments
- tutorials
- practical suggestions

STEP 1:
Analyze weaknesses in current syllabus.

STEP 2:
Determine what students SHOULD ideally learn globally.

STEP 3:
Redesign prerequisites logically.

STEP 4:
Redesign course objectives intelligently.

STEP 5:
Generate ideal course outcomes.

IMPORTANT:
For each CO:
- identify Bloom level
- justify Bloom level
- explain expected competency

STEP 6 — HOLISTIC INTELLIGENT CO-PO-PSO MAPPING

IMPORTANT:

Do NOT generate random mappings.

You must THINK deeply like:
- NBA evaluator
- accreditation committee member
- curriculum expert
- OBE strategist

IMPORTANT:

PO mapping must NOT be decided
only from technical content.

PO mapping must ALSO consider:

- pedagogy strategy
- student activities
- collaborative learning
- seminar activities
- peer teaching
- communication opportunities
- self-learning opportunities
- project execution
- assignment structure
- inquiry-based learning
- presentation opportunities
- teamwork exposure
- problem-solving workshops
- experiential learning

IMPORTANT:

Even analytical/theoretical subjects
may contribute to:
- communication skills
- teamwork
- ethics
- project management
- lifelong learning

IF pedagogy and activities intentionally support them.

Therefore:

Mapping logic must consider BOTH:
1. Technical competency contribution
2. Learning-process competency contribution

------------------------------------------------

For each CO:

Determine:
- which PO is genuinely supported
- which PSO is genuinely supported
- what mapping strength is realistic

Use mapping levels:
- 1 = Low
- 2 = Moderate
- 3 = Strong

IMPORTANT:

Mapping strength must be based on:
- actual competency contribution
- Bloom level
- practical exposure
- analytical depth
- communication exposure
- teamwork exposure
- tool usage
- problem solving
- real-world application
- self-learning exposure
- presentation exposure
- collaborative learning exposure
- inquiry-based learning exposure

IMPORTANT:

The AI MUST intelligently map suitable COs
with:
- communication PO
- ethics PO
- teamwork PO
- project management PO
- lifelong learning PO

WHEN pedagogy/activity/assessment supports them.

------------------------------------------------

VERY IMPORTANT:

After generating mapping matrix,
generate INDIVIDUAL JUSTIFICATION
for EVERY SINGLE mapping value.

FORMAT STRICTLY LIKE THIS:

MAPPING OF CO1 TO PO1 (VALUE 2):
Explain:
- WHY this mapping exists
- HOW CO1 contributes to PO1
- WHY contribution strength is MODERATE
- What competency dimension supports this mapping

MAPPING OF CO1 TO PO9 (VALUE 2):
Explain:
- HOW collaborative pedagogy supports teamwork
- HOW group analytical activities support PO9
- WHY teamwork contribution is MODERATE

MAPPING OF CO1 TO PO10 (VALUE 1):
Explain:
- HOW presentation/discussion/self-expression
supports communication competency
- WHY communication contribution is LOW

MAPPING OF CO1 TO PO12 (VALUE 2):
Explain:
- HOW inquiry/self-learning/open-resource exploration
supports lifelong learning competency

IMPORTANT:

DO THIS FOR:
- EVERY CO
- EVERY mapped PO
- EVERY mapped PSO

The justification must NOT be generic.

Each justification must be:
- competency-specific
- Bloom-specific
- logically defendable
- accreditation-ready
- academically meaningful

The explanation must clearly show:
- cognitive contribution
- analytical contribution
- practical contribution
- communication contribution
- teamwork contribution
- tool contribution
- engineering contribution
- real-world application contribution
- lifelong learning contribution
- project management contribution

IMPORTANT:

The justification should be detailed enough
that an accreditation evaluator can clearly understand:
WHY the mapping value is logically correct.

STEP 7:
Generate UNIT-WISE pedagogy.

For each unit:
- pedagogy method
- exact activity
- how activity should be conducted
- mapped CO
- mapped PO
- mapped PSO
- graduate competency developed

Use:
- Inquiry Based Learning
- Peer Teaching
- Group Learning
- Problem Solving
- Technology Based Learning
- Game Based Learning
- Collaborative Learning
- Any other Student Centric Teaching Learning Activities 
STEP 8:
Generate UNIT-WISE ICT tools.

STEP 9:
Generate 3 HIGH VALUE TERM WORKS.

Assignments must:
- cover all units
- support CO attainment
- support PO attainment
- improve employability
- improve communication
- improve teamwork
- improve practical competency

For each assignment provide:
- title
- objective
- execution strategy
- mapped COs
- mapped POs
- mapped PSOs
- Bloom level

STEP 10:
Suggest improved practicals and tutorials.

STEP 11:
Generate industry-oriented activities.

STEP 12:
Generate SDG mapping.

STEP 13:
Generate final modern accreditation-ready course file.

VISION & MISSION:
{vision_text}

PEO PO PSO:
{peo_text}

COURSE FILE:
{course_text}

"""

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[
            {
                "role": "system",
                "content": "You are a globally experienced curriculum intelligence expert."
            },
            {
                "role": "user",
                "content": prompt
            }
        ],
        temperature=0.4
    )

    return response.choices[0].message.content

def create_docx(content):

    document = Document()

    document.add_heading(
        "AI Generated Curriculum Intelligence Report",
        level=1
    )

    paragraphs = content.split("\n")

    for para in paragraphs:

        para = para.strip()

        if para:

            if len(para) < 120 and (
                para.isupper() or
                para.endswith(":")
            ):

                document.add_heading(para, level=2)

            else:

                document.add_paragraph(para)

    buffer = BytesIO()

    document.save(buffer)

    buffer.seek(0)

    return buffer

if (
    vision_file is not None and
    peo_file is not None and
    course_file is not None
):

    try:

        with st.spinner("Reading uploaded documents..."):

            vision_text = extract_text_from_pdf(vision_file)

            peo_text = extract_text_from_pdf(peo_file)

            course_text = extract_text_from_pdf(course_file)

        with st.spinner("AI is redesigning curriculum intelligently..."):

            refined_output = generate_curriculum_intelligence_report(
                vision_text,
                peo_text,
                course_text
            )

        st.success(
            "AI Curriculum Intelligence Report Generated Successfully"
        )

        st.markdown(refined_output)

        docx_file = create_docx(refined_output)

        st.download_button(
            label="Download Editable DOCX Report",
            data=docx_file,
            file_name="AI_Curriculum_Intelligence_Report.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    except Exception as e:

        st.error(f"Error: {e}")
